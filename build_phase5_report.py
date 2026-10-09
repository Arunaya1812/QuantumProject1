from pathlib import Path
import json

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "outputs" / "Phase5_Improvement_Report.docx"
REPORT = json.loads((ROOT / "outputs" / "phase5_evaluation_report.json").read_text())
LITERATURE = pd.read_csv(ROOT / "outputs" / "phase5_literature_comparison.csv").fillna("")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, widths=None, font_size=9):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_index, row in enumerate(table.rows):
        for c_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = widths[c_index]
            if r_index == 0:
                set_cell_shading(cell, "1F4E78")
            elif r_index % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    if r_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(9)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def pct(value):
    return f"{100 * value:.2f}%"


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Normal" else "Aptos"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)
    style_ppr = style.element.get_or_add_pPr()
    style_border = style_ppr.find(qn("w:pBdr"))
    if style_border is not None:
        style_ppr.remove(style_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Phase 5 Unified Adaptive ML Evaluation")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run(
    "Adaptive ML Based Eavesdropper Detection in Decoy State BB84 QKD with Post Quantum Key Authentication"
).bold = True
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Evaluation date 10 October 2026\n").italic = True
meta.add_run("Prepared for project review")

doc.add_heading("Executive Result", level=1)
unified = REPORT["metrics"]["unified"]
ml = REPORT["metrics"]["ml"]
adaptive = REPORT["metrics"]["adaptive"]
static = REPORT["metrics"]["static"]
doc.add_paragraph(
    "The Phase 5 detector is the strongest method among the four project baselines on the same "
    "untouched test timelines. It improves on the standalone Random Forest in accuracy, precision, "
    "recall, F1 score, specificity, false-positive rate, and false-negative rate. The decoy-aware "
    "adaptive baseline also improves substantially over the conventional fixed 11 percent QBER rule."
)
summary = doc.add_table(rows=1, cols=5)
for cell, value in zip(summary.rows[0].cells, ["Detector", "Accuracy", "Precision", "Recall", "F1"]):
    cell.text = value
for name in ("static", "adaptive", "ml", "unified"):
    values = REPORT["metrics"][name]
    cells = summary.add_row().cells
    cells[0].text = name.title()
    cells[1].text = pct(values["accuracy"])
    cells[2].text = pct(values["precision"])
    cells[3].text = pct(values["recall_detection_rate"])
    cells[4].text = pct(values["f1_score"])
style_table(summary, [Inches(1.35)] * 5, 9)

doc.add_heading("Phase 5 Architecture", level=1)
doc.add_paragraph(
    "The original unified detector used a simple logical OR between the adaptive QBER alarm and the "
    "Random Forest decision. That raised recall but could also inherit false positives from either "
    "component. Phase 5 replaces it with a stacked, confidence-aware fusion layer."
)
for item in REPORT["novelty"]:
    doc.add_paragraph(item, style="List Bullet")
doc.add_paragraph(
    "The Random Forest remains the primary classifier. A lightweight histogram-gradient-boosting "
    "fusion model receives the RF attack probability, its three-window mean and maximum, the change "
    "in probability, calibration-normalized QBER and decoy-gap z-scores, and the two adaptive flags. "
    "A probability guard prevents suspicious windows from updating the adaptive baselines."
)

doc.add_heading("Leakage Safe Evaluation", level=1)
design = REPORT["evaluation_design"]
doc.add_paragraph(
    f"The dataset contains {design['total_timelines']} independent timelines. "
    f"Seventy timelines were used for development and {design['untouched_test_timelines']} complete "
    f"timelines containing {design['test_windows']} windows were kept untouched for final evaluation. "
    "Windows from the same timeline were never divided between development and test data."
)
doc.add_paragraph(
    "The fusion layer was trained from five-fold out-of-fold Random Forest probabilities. Therefore, "
    "the fusion model never learned from a base-model probability generated on the same timeline used "
    "to train that base model. This reduces optimistic leakage in a temporal dataset."
)

doc.add_picture(str(ROOT / "outputs" / "phase5_detector_metrics.png"), width=Inches(7.0))
add_caption(doc, "Figure 1 Held out comparison of fixed adaptive ML and unified detectors")

doc.add_heading("Expanded Performance Metrics", level=1)
fields = [
    ("balanced_accuracy", "Balanced accuracy"),
    ("specificity", "Specificity"),
    ("f2_score", "F2 score"),
    ("matthews_correlation_coefficient", "MCC"),
    ("false_positive_rate", "FPR"),
    ("false_negative_rate", "FNR"),
    ("roc_auc", "ROC AUC"),
    ("pr_auc", "PR AUC"),
]
expanded = doc.add_table(rows=1, cols=5)
for cell, value in zip(expanded.rows[0].cells, ["Metric", "Static", "Adaptive", "ML", "Unified"]):
    cell.text = value
for key, label in fields:
    cells = expanded.add_row().cells
    cells[0].text = label
    for index, name in enumerate(("static", "adaptive", "ml", "unified"), start=1):
        value = REPORT["metrics"][name][key]
        cells[index].text = f"{value:.4f}"
style_table(expanded, [Inches(1.65), Inches(1.25), Inches(1.25), Inches(1.25), Inches(1.25)], 8.5)

doc.add_picture(str(ROOT / "outputs" / "phase5_confusion_matrices.png"), width=Inches(6.8))
add_caption(doc, "Figure 2 Confusion matrices on 2160 windows from untouched test timelines")

doc.add_heading("What Improved", level=2)
doc.add_paragraph(
    f"Adaptive versus static: recall rises from {pct(static['recall_detection_rate'])} to "
    f"{pct(adaptive['recall_detection_rate'])}, and F1 rises from {pct(static['f1_score'])} to "
    f"{pct(adaptive['f1_score'])}. This happens because the adaptive baseline now monitors the "
    "signal-decoy gap as well as QBER, allowing it to detect PNS behavior that a fixed QBER rule "
    "cannot observe."
)
doc.add_paragraph(
    f"Unified versus ML: accuracy rises from {pct(ml['accuracy'])} to {pct(unified['accuracy'])}; "
    f"precision rises from {pct(ml['precision'])} to {pct(unified['precision'])}; recall rises from "
    f"{pct(ml['recall_detection_rate'])} to {pct(unified['recall_detection_rate'])}; and F1 rises "
    f"from {pct(ml['f1_score'])} to {pct(unified['f1_score'])}. False positives fall from "
    f"{ml['confusion_matrix']['fp']} to {unified['confusion_matrix']['fp']}, while false negatives "
    f"fall from {ml['confusion_matrix']['fn']} to {unified['confusion_matrix']['fn']}."
)
test_result = REPORT["paired_mcnemar_unified_vs_ml"]
doc.add_paragraph(
    f"The paired exact McNemar p-value is {test_result['p_value']:.3f}. The measured improvement is "
    "consistent across the reported metrics, but it is not statistically significant at the 0.05 "
    "level on this 18-timeline test set. More independent timelines are required for a stronger "
    "superiority claim."
)

doc.add_heading("Five Paper Literature Comparison", level=1)
doc.add_paragraph(
    "The following five papers were selected from the Review 1 literature survey. Numeric values are "
    "shown only when the paper reports a comparable classification metric. Different protocols, "
    "datasets, attack sets, and split designs prevent a formal head-to-head ranking."
)
lit_table = doc.add_table(rows=1, cols=5)
headers = ["Paper", "Task", "Reported result", "Project comparison", "Interpretation"]
for cell, value in zip(lit_table.rows[0].cells, headers):
    cell.text = value
for _, row in LITERATURE.iterrows():
    reported = []
    for label, key in (("Acc", "reported_accuracy"), ("Prec", "reported_precision"),
                       ("Recall", "reported_recall"), ("F1", "reported_f1")):
        value = row[key]
        if value != "":
            if isinstance(value, (float, int)):
                reported.append(f"{label} {float(value):.4f}")
            else:
                reported.append(f"{label} {value}")
    cells = lit_table.add_row().cells
    cells[0].text = str(row["reference"])
    cells[1].text = str(row["protocol_task"])
    cells[2].text = "; ".join(reported) or "No classification metric"
    cells[3].text = str(row["project_comparable_metric"])
    cells[4].text = str(row["comparability"])
style_table(
    lit_table,
    [Inches(1.25), Inches(1.35), Inches(1.05), Inches(1.55), Inches(2.0)],
    7.4,
)

doc.add_heading("Defensible Novelty Claim", level=1)
doc.add_paragraph(
    "The defensible novelty is not that one classifier universally defeats every published method. "
    "It is the integration of decoy-aware dual adaptive baselines, baseline-poisoning protection, "
    "temporal RF confidence, stacked fusion trained without timeline leakage, Qiskit refinement, and "
    "ML-DSA plus ML-KEM authenticated hybrid key establishment in one end-to-end prototype."
)
doc.add_paragraph(
    "Recommended review statement: On the same held-out simulated timelines, our unified detector "
    "achieved the best measured balance of accuracy, precision, recall, F1, specificity, FPR and FNR "
    "among our fixed, adaptive, and standalone RF baselines. Its numeric results also exceed the "
    "reported values of Papers 6, 7 and 8 on selected metrics, but those cross-paper figures are "
    "contextual because the experimental datasets differ."
)

doc.add_heading("Limitations", level=1)
for limitation in [
    "All detector results use simulated data; physical optical hardware has not been evaluated.",
    "The final test contains 18 timelines, and the ML versus unified difference is not yet significant at p < 0.05.",
    "Cross-paper values are not directly comparable without a shared dataset and attack definition.",
    "The Qiskit backend models BB84 qubit gates, while weak coherent optical decoy pulses remain statistical.",
    "Error correction is represented at protocol level and is not a complete Cascade or LDPC implementation.",
]:
    doc.add_paragraph(limitation, style="List Bullet")

doc.add_heading("Review Commands", level=1)
commands = [
    "python train_phase5_unified.py",
    "python compare_detectors.py",
    "Get-Content .\\outputs\\phase5_evaluation_report.json",
    "Invoke-Item .\\outputs\\phase5_detector_metrics.png",
    "Invoke-Item .\\outputs\\phase5_confusion_matrices.png",
    "python run_phase3_demo.py --attack none --windows 40 --photons 1000 --seed 2026",
    "python run_phase3_demo.py --attack pns --attack-start 0.75 --windows 40 --seed 2026",
]
for command in commands:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(command)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)

doc.add_heading("References", level=1)
for _, row in LITERATURE.iterrows():
    paragraph = doc.add_paragraph(f"{row['reference']}. {row['source']}")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0

doc.save(OUTPUT)
print(OUTPUT)
